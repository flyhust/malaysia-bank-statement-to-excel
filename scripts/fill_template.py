#!/usr/bin/env python3
"""Fill an Excel template from verified statement JSON, then read it back and re-check.

Usage:
  python fill_template.py statement.json out.xlsx [--template tpl.xlsx]
  python fill_template.py months.json out_dir/ [--template tpl.xlsx]     # one file per statement, e.g. JAN_2025.xlsx
  python fill_template.py months.json all.xlsx --combine [--template tpl.xlsx]

Without --template a default template is built (see references/data-format.md).
Requires openpyxl.
"""
import argparse
import datetime as dt
import json
import os
import sys
from copy import copy
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from verify import money, verify_all  # noqa: E402

STATEMENT_FIELDS = [
    ("Period start", "period_start"),
    ("Period end", "period_end"),
    ("Holder name", "holder_name"),
    ("Holder address", "holder_address"),
    ("Account name", "account_name"),
    ("Account number", "account_number"),
    ("Opening balance", "opening_balance"),
    ("Ending balance", "closing_balance"),
]
TX_COLUMNS = [
    ("Date", "date"),
    ("Description", "description"),
    ("Cheque / serial no", "cheque_no"),
    ("Debit", "debit"),
    ("Credit", "credit"),
    ("Balance", "balance"),
    ("Counterparty", "counterparty"),
    ("Category", "category"),
    ("Account code", "account_code"),
]
HANDWRITTEN = "Invoice no (handwritten)"
MONEY_KEYS = {"debit", "credit", "balance", "opening_balance", "closing_balance"}
# Header aliases so a user's template with different wording still maps.
ALIASES = {
    "date": "date", "transaction date": "date",
    "description": "description", "details": "description", "narrative": "description",
    "cheque / serial no": "cheque_no", "cheque no": "cheque_no", "serial no": "cheque_no", "reference": "cheque_no",
    "debit": "debit", "withdrawal": "debit", "dr": "debit",
    "credit": "credit", "deposit": "credit", "cr": "credit",
    "balance": "balance",
    "counterparty": "counterparty", "payee / payer": "counterparty", "name": "counterparty",
    "category": "category", "account code": "account_code",
    HANDWRITTEN.lower(): "handwritten",
}
FIELD_ALIASES = {label.lower(): key for label, key in STATEMENT_FIELDS}
FIELD_ALIASES.update({"closing balance": "closing_balance", "ending balance": "closing_balance"})


def norm(s):
    return " ".join(str(s or "").strip().lower().split())


def default_workbook():
    wb = Workbook()
    font = Font(name="Calibri", size=12)
    ws = wb.active
    ws.title = "Statement"
    ws.append(["Field", "Value"])
    for label, _ in STATEMENT_FIELDS:
        ws.append([label, None])
    tx = wb.create_sheet("Transactions")
    tx.append([h for h, _ in TX_COLUMNS])
    for sheet in (ws, tx):
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = Font(name="Calibri", size=12, bold=cell.row == 1)
    return wb, font


def find_sheet(wb, name, probe):
    for ws in wb.worksheets:
        if norm(ws.title) == name:
            return ws
    for ws in wb.worksheets:
        if probe(ws):
            return ws
    raise SystemExit(f"template has no '{name}' sheet")


def header_row(ws):
    for row in ws.iter_rows(min_row=1, max_row=10):
        names = [norm(c.value) for c in row]
        if "date" in names or "transaction date" in names:
            return row[0].row
    raise SystemExit(f"no header row with 'Date' found in sheet '{ws.title}'")


def combine(statements):
    first, last = statements[0], statements[-1]
    out = {k: v for k, v in last.items() if k not in ("rows", "summary")}
    out["opening_balance"] = first["opening_balance"]
    out["period_start"] = first.get("period_start")
    out["rows"] = sorted((r for s in statements for r in s.get("rows") or []), key=lambda r: r.get("date") or "")
    return out


def write(st, path, template):
    if template:
        wb = load_workbook(template)
        font = None
    else:
        wb, font = default_workbook()

    sws = find_sheet(wb, "statement", lambda w: norm(w.cell(1, 1).value) == "field")
    for row in sws.iter_rows(min_row=1):
        key = FIELD_ALIASES.get(norm(row[0].value))
        if key and len(row) > 1:
            v = st.get(key)
            row[1].value = money(v, key) if key in MONEY_KEYS else v

    tws = find_sheet(wb, "transactions", lambda w: True)
    h = header_row(tws)
    cols = {}
    for cell in tws[h]:
        key = ALIASES.get(norm(cell.value))
        if key:
            cols[key] = cell.column
    rows = st.get("rows") or []
    if any(r.get("handwritten") for r in rows) and "handwritten" not in cols:
        col = tws.max_column + 1
        hc = tws.cell(h, col, HANDWRITTEN)
        hc.font = copy(tws.cell(h, 1).font)
        cols["handwritten"] = col

    # Style of the first data row if the template has one, else header font unbolded.
    styles = {}
    for key, col in cols.items():
        sample = tws.cell(h + 1, col)
        f = copy(sample.font) if sample.has_style else copy(tws.cell(h, col).font)
        if font is None and not sample.has_style:
            f = Font(name=f.name, size=f.size)
        styles[key] = (font or f, sample.number_format if sample.has_style else None)

    if tws.max_row > h:
        tws.delete_rows(h + 1, tws.max_row - h)
    for i, r in enumerate(rows, h + 1):
        for key, col in cols.items():
            v = r.get(key)
            if key in MONEY_KEYS:
                v = money(v, key)
            cell = tws.cell(i, col, v)
            f, fmt = styles[key]
            cell.font = copy(f)
            if key in MONEY_KEYS and v is not None:
                cell.number_format = fmt if fmt and fmt != "General" else "#,##0.00"
    wb.save(path)


def read_back(path):
    """Rebuild statement JSON from the saved file so verify sees what was written."""
    wb = load_workbook(path, data_only=True)
    sws = find_sheet(wb, "statement", lambda w: norm(w.cell(1, 1).value) == "field")
    st = {}
    for row in sws.iter_rows(min_row=1):
        key = FIELD_ALIASES.get(norm(row[0].value))
        if key and len(row) > 1:
            st[key] = row[1].value
    tws = find_sheet(wb, "transactions", lambda w: True)
    h = header_row(tws)
    cols = {ALIASES.get(norm(c.value)): c.column for c in tws[h] if ALIASES.get(norm(c.value))}
    st["rows"] = []
    for i in range(h + 1, tws.max_row + 1):
        r = {k: tws.cell(i, c).value for k, c in cols.items()}
        if all(v is None for v in r.values()):
            continue
        for k in MONEY_KEYS & r.keys():
            r[k] = None if r[k] is None else str(Decimal(str(r[k])))
        st["rows"].append(r)
    return st


def month_name(st):
    d = st.get("statement_date") or st.get("period_end") or ""
    try:
        return dt.date.fromisoformat(d).strftime("%b_%Y").upper()
    except ValueError:
        return d or "statement"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("json")
    ap.add_argument("out")
    ap.add_argument("--template")
    ap.add_argument("--combine", action="store_true")
    a = ap.parse_args()

    with open(a.json, encoding="utf-8") as f:
        data = json.load(f)
    statements = data if isinstance(data, list) else [data]

    errs = verify_all(statements)
    if errs:
        print("Refusing to write: the transcription does not reconcile.")
        for e in errs:
            print("  - " + e)
        return 1

    if a.combine:
        jobs = [(combine(statements), a.out)]
    elif len(statements) == 1 and a.out.lower().endswith(".xlsx"):
        jobs = [(statements[0], a.out)]
    else:
        os.makedirs(a.out, exist_ok=True)
        jobs = [(s, os.path.join(a.out, month_name(s) + ".xlsx")) for s in statements]

    failed = False
    for st, path in jobs:
        write(st, path, a.template)
        back = read_back(path)
        back.setdefault("summary", None)
        errs = verify_all(back)
        if errs:
            failed = True
            print(f"{path}: written but read-back check FAILED")
            for e in errs:
                print("  - " + e)
        else:
            print(f"{path}: written, read back, re-checked: {len(back['rows'])} row(s) reconcile")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
